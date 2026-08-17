"""Pure file loaders for distributions (`.csv` / `.xlsx` / `.json`).

Kept **inside** datagen (stdlib ``csv`` + optional ``openpyxl`` only) so the
package stays isolated — it imports nothing from the scenario layer. The scenario
layer's ``scenario_inputs`` resolves ``{"$file": path}`` at spec-creation and
inlines the result, so at compile time a distribution usually already has an
inline dict; ``load_*_file`` is used when a policy names a data file directly.
"""

from __future__ import annotations

import csv
import json
import os
from typing import Any

_SUPPORTED = (".json", ".csv", ".xlsx")
_HOUR_ALIASES = {"hour", "h", "t"}
_WEIGHT_ALIASES = {"weight", "w", "value", "count", "pct"}


def _ext(path: str) -> str:
    return os.path.splitext(path)[1].lower()


def _read_rows(path: str) -> list[list[Any]]:
    ext = _ext(path)
    if ext == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as fp:
            return [list(row) for row in csv.reader(fp)]
    if ext == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover
            raise ValueError("Reading .xlsx requires the 'openpyxl' package.") from exc
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
        return rows
    raise ValueError(f"Unsupported spreadsheet extension {ext!r} (use .csv or .xlsx)")


def _cell(v: Any) -> str:
    return "" if v is None else str(v).strip()


def load_matrix_file(path: str) -> dict[str, dict[str, float]]:
    """Parse a `{row: {col: weight}}` grid from `.json`/`.csv`/`.xlsx`."""
    if not os.path.isfile(path):
        raise ValueError(f"matrix file not found: {path}")
    if _ext(path) == ".json":
        with open(path, encoding="utf-8") as fp:
            data = json.load(fp)
        if not isinstance(data, dict):
            raise ValueError("Matrix JSON must be an object of {row: {col: weight}}.")
        return data
    rows = _read_rows(path)
    if len(rows) < 2:
        raise ValueError("Matrix needs a header row and at least one data row.")
    header = [_cell(c).upper() for c in rows[0]]
    col_codes = header[1:]
    grid: dict[str, dict[str, float]] = {}
    for row in rows[1:]:
        rk = _cell(row[0]).upper() if row else ""
        if not rk:
            continue
        cells = grid.setdefault(rk, {})
        for i, ck in enumerate(col_codes):
            if not ck:
                continue
            try:
                cells[ck] = max(0.0, float(row[i + 1]))
            except (IndexError, TypeError, ValueError):
                cells[ck] = 0.0
    if not grid:
        raise ValueError("No matrix rows found.")
    return grid


def load_curve_file(path: str) -> dict[str, Any]:
    """Parse a `{resolution:'hour', points:[{hour,weight}]}` curve from a file."""
    if not os.path.isfile(path):
        raise ValueError(f"curve file not found: {path}")
    if _ext(path) == ".json":
        with open(path, encoding="utf-8") as fp:
            return json.load(fp)
    rows = _read_rows(path)
    if not rows:
        raise ValueError("Curve file is empty.")
    header = [_cell(c).lower() for c in rows[0]]
    hour_col = next((i for i, c in enumerate(header) if c in _HOUR_ALIASES), -1)
    weight_col = next((i for i, c in enumerate(header) if c in _WEIGHT_ALIASES), -1)
    if hour_col < 0 or weight_col < 0:
        raise ValueError('Curve file needs "hour" and "weight" columns.')
    points: list[dict[str, float]] = []
    for row in rows[1:]:
        try:
            hour = int(round(float(row[hour_col])))
        except (IndexError, TypeError, ValueError):
            continue
        if not (0 <= hour <= 23):
            continue
        try:
            weight = max(0.0, float(row[weight_col]))
        except (IndexError, TypeError, ValueError):
            weight = 0.0
        points.append({"hour": hour, "weight": weight})
    if not points:
        raise ValueError("No valid hour rows (0-23) in curve file.")
    return {"resolution": "hour", "points": points}


def load_records_file(path: str) -> list[dict]:
    """Parse historical records (list of dicts) from `.json`/`.csv`/`.xlsx`."""
    if not os.path.isfile(path):
        raise ValueError(f"records file not found: {path}")
    if _ext(path) == ".json":
        with open(path, encoding="utf-8") as fp:
            data = json.load(fp)
        if isinstance(data, list):
            return [r for r in data if isinstance(r, dict)]
        raise ValueError("Records JSON must be a list of objects.")
    rows = _read_rows(path)
    if len(rows) < 2:
        raise ValueError("Records file needs a header row and at least one data row.")
    header = [_cell(c) for c in rows[0]]
    out: list[dict] = []
    for row in rows[1:]:
        rec = {header[i]: (row[i] if i < len(row) else None) for i in range(len(header)) if header[i]}
        if any(v is not None and _cell(v) for v in rec.values()):
            out.append(rec)
    return out

"""Parse uploaded trip-matrix / order-demand-curve files into spec values.

Used at the **spec-creation** step (``normalize_generate_spec``): a spec field given as
``{"$file": "/abs/path.csv"}`` is parsed here and the result is **inlined** into
``spec.json`` — no file is referenced at compile or run time. Supports ``.json``,
``.csv`` and ``.xlsx`` (Excel via ``openpyxl``), in the same layout as the dashboard's
client-side importer (``analytics/lib/matrixImport.ts``) so every entry point agrees.

The returned structures are the *raw* shapes the existing normalizers consume
(``datagen.trip_matrix.parse_trip_matrix`` for the grid;
``order_demand.parse_order_demand_curve`` for the curve) — normalization stays in one
place in ``normalize_generate_spec``.
"""

from __future__ import annotations

import csv
import json
import os
from typing import Any

MATRIX_FILE_KEY = "$file"
_SUPPORTED_EXT = (".json", ".csv", ".xlsx")

# Demand-curve column header aliases (mirror matrixImport.ts).
_HOUR_ALIASES = {"hour", "h", "t"}
_WEIGHT_ALIASES = {"weight", "w", "value", "count", "pct"}

_TEMPLATE_HINT = "Download the template for the exact layout."


def _ext(path: str) -> str:
    return os.path.splitext(path)[1].lower()


def _read_rows(path: str) -> list[list[Any]]:
    """Read a .csv or .xlsx file as a list of row-lists (first sheet for xlsx)."""
    ext = _ext(path)
    if ext == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as fp:
            return [list(row) for row in csv.reader(fp)]
    if ext == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover - dependency guard
            raise ValueError(
                "Reading .xlsx requires the 'openpyxl' package in the apps venv."
            ) from exc
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
        return rows
    raise ValueError(f"Unsupported spreadsheet extension {ext!r} (use .csv or .xlsx)")


def _cell(v: Any) -> str:
    return "" if v is None else str(v).strip()


def parse_trip_matrix_file(path: str) -> dict[str, dict[str, float]]:
    """Parse a trip-matrix file into a ``{pickup: {delivery: weight}}`` grid (raw).

    Layout (csv/xlsx): first row ``[corner, CT, CU, MT, …]`` (delivery codes), each data
    row ``[pickup, w_CT, w_CU, …]``. ``.json`` is loaded as the grid directly. Codes are
    upper-cased; downstream ``parse_trip_matrix`` forces the diagonal to 0 and normalizes.
    """
    if not os.path.isfile(path):
        raise ValueError(f"file not found: {path}")
    if _ext(path) == ".json":
        with open(path, "r", encoding="utf-8") as fp:
            data = json.load(fp)
        if not isinstance(data, dict):
            raise ValueError("Trip-matrix JSON must be an object of {pickup: {delivery: weight}}.")
        return data

    rows = _read_rows(path)
    if len(rows) < 2:
        raise ValueError(f"Trip matrix needs a header row and at least one data row. {_TEMPLATE_HINT}")
    header = [_cell(c).upper() for c in rows[0]]
    col_codes = header[1:]
    grid: dict[str, dict[str, float]] = {}
    for row in rows[1:]:
        pickup = _cell(row[0]).upper() if row else ""
        if not pickup:
            continue
        cells = grid.setdefault(pickup, {})
        for i, delivery in enumerate(col_codes):
            if not delivery:
                continue
            try:
                w = float(row[i + 1])
            except (IndexError, TypeError, ValueError):
                w = 0.0
            cells[delivery] = max(0.0, w)
    if not grid:
        raise ValueError(f"No matrix rows found. {_TEMPLATE_HINT}")
    return grid


def parse_demand_curve_file(path: str) -> dict[str, Any]:
    """Parse an order-demand-curve file into ``{resolution: 'hour', points: [...]}`` (raw).

    Layout (csv/xlsx): a header row with ``hour`` and ``weight`` columns (aliases allowed),
    then rows for hours 0–23. ``.json`` is loaded directly. Downstream
    ``parse_order_demand_curve`` interpolates gaps and normalizes to 24 weights summing to 1.
    """
    if not os.path.isfile(path):
        raise ValueError(f"file not found: {path}")
    if _ext(path) == ".json":
        with open(path, "r", encoding="utf-8") as fp:
            return json.load(fp)

    rows = _read_rows(path)
    if not rows:
        raise ValueError(f"Demand-curve sheet is empty. {_TEMPLATE_HINT}")
    header = [_cell(c).lower() for c in rows[0]]
    hour_col = next((i for i, c in enumerate(header) if c in _HOUR_ALIASES), -1)
    weight_col = next((i for i, c in enumerate(header) if c in _WEIGHT_ALIASES), -1)
    if hour_col < 0 or weight_col < 0:
        raise ValueError(f'Expected "hour" and "weight" columns in the first row. {_TEMPLATE_HINT}')
    points: list[dict[str, float]] = []
    for row in rows[1:]:
        try:
            hour = int(round(float(row[hour_col])))
        except (IndexError, TypeError, ValueError):
            continue
        if not (0 <= hour <= 23):
            continue
        try:
            weight = max(0.0, float(row[weight_col]))
        except (IndexError, TypeError, ValueError):
            weight = 0.0
        points.append({"hour": hour, "weight": weight})
    if not points:
        raise ValueError(f"No valid hour rows (0–23) found. {_TEMPLATE_HINT}")
    return {"resolution": "hour", "points": points}


def maybe_parse_spec_file(value: Any, kind: str) -> Any:
    """If ``value`` is ``{"$file": path}``, parse the file by ``kind`` ('matrix'|'curve');
    otherwise return ``value`` unchanged (inline dicts pass straight through)."""
    if not (isinstance(value, dict) and MATRIX_FILE_KEY in value):
        return value
    path = value.get(MATRIX_FILE_KEY)
    if not isinstance(path, str) or not path.strip():
        raise ValueError(f"{MATRIX_FILE_KEY} must be a non-empty file path")
    path = path.strip()
    if _ext(path) not in _SUPPORTED_EXT:
        raise ValueError(f"Unsupported file type {_ext(path)!r} — use .json, .csv, or .xlsx")
    if kind == "matrix":
        return parse_trip_matrix_file(path)
    if kind == "curve":
        return parse_demand_curve_file(path)
    raise ValueError(f"Unknown spec-file kind: {kind!r}")
